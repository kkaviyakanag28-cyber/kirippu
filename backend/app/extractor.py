import os
import io
import re
import logging
from typing import List, Dict, Any, Tuple

logger = logging.getLogger(__name__)


class TextExtractor:
    """Extracts raw text and metadata from PDF, DOCX, TXT, and Image files."""

    @classmethod
    async def extract(cls, file_bytes: bytes, filename: str, mime_type: str) -> Dict[str, Any]:
        ext = os.path.splitext(filename)[1].lower()
        pages: List[Dict[str, Any]] = []

        try:
            if ext == ".pdf" or "pdf" in mime_type:
                pages = cls._extract_pdf(file_bytes)
            elif ext in [".docx", ".doc"] or "wordprocessingml" in mime_type or "msword" in mime_type:
                pages = cls._extract_docx(file_bytes)
            elif ext in [".txt", ".md", ".csv", ".json"] or "text/" in mime_type:
                pages = cls._extract_txt(file_bytes)
            elif ext in [".png", ".jpg", ".jpeg", ".webp"] or "image/" in mime_type:
                pages = cls._extract_image(file_bytes)
            else:
                # Fallback to text decoding
                pages = cls._extract_txt(file_bytes)
        except Exception as e:
            logger.error(f"Error during extraction for {filename}: {e}")
            pages = [{"page": 1, "text": f"Error extracting content: {str(e)}"}]

        full_text = "\n\n".join([p["text"] for p in pages if p.get("text")])
        words = full_text.split()
        
        return {
            "pages": pages,
            "full_text": full_text,
            "word_count": len(words),
            "page_count": len(pages),
        }

    @classmethod
    def _extract_pdf(cls, file_bytes: bytes) -> List[Dict[str, Any]]:
        import PyPDF2
        pages = []
        stream = io.BytesIO(file_bytes)
        reader = PyPDF2.PdfReader(stream)
        for i, page in enumerate(reader.pages):
            text = page.extract_text() or ""
            pages.append({"page": i + 1, "text": text.strip()})
        return pages if pages else [{"page": 1, "text": ""}]

    @classmethod
    def _extract_docx(cls, file_bytes: bytes) -> List[Dict[str, Any]]:
        import docx
        stream = io.BytesIO(file_bytes)
        doc = docx.Document(stream)
        text_lines = []
        for p in doc.paragraphs:
            if p.text.strip():
                text_lines.append(p.text.strip())
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join([cell.text.strip() for cell in row.cells if cell.text.strip()])
                if row_text:
                    text_lines.append(row_text)
        return [{"page": 1, "text": "\n".join(text_lines)}]

    @classmethod
    def _extract_txt(cls, file_bytes: bytes) -> List[Dict[str, Any]]:
        for enc in ["utf-8", "utf-8-sig", "latin-1", "cp1252"]:
            try:
                text = file_bytes.decode(enc)
                return [{"page": 1, "text": text.strip()}]
            except UnicodeDecodeError:
                continue
        return [{"page": 1, "text": file_bytes.decode("utf-8", errors="ignore")}]

    @classmethod
    def _extract_image(cls, file_bytes: bytes) -> List[Dict[str, Any]]:
        try:
            from PIL import Image
            import pytesseract
            img = Image.open(io.BytesIO(file_bytes))
            text = pytesseract.image_to_string(img)
            if text.strip():
                return [{"page": 1, "text": text.strip()}]
            return [{"page": 1, "text": "[Image uploaded: No readable text detected with OCR.]"}]
        except Exception as e:
            logger.info(f"Tesseract OCR not installed or failed: {e}. Providing friendly fallback.")
            return [{
                "page": 1,
                "text": "[Image uploaded: Optical Character Recognition (OCR) engine not configured locally. Upload text, PDF, or DOCX for full automated action extraction.]"
            }]
